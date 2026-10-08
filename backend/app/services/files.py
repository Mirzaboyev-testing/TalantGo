"""Files students attach to their work (Excel, PDF, screenshots, notebooks...).

Stored under a random name on the server's disk and served only through signed links that expire, so a
file is reachable only by people the app showed it to (the author, reviewers on the track, mentors).
Anything a browser could run (HTML, SVG...) is served as a download, never rendered on our domain.
"""
import base64
import json
import mimetypes
import re
import uuid
from pathlib import Path
from urllib.parse import quote

from fastapi import HTTPException, UploadFile

from ..core.config import settings
from ..core.security import create_token
from ..models import SubmissionFile

MAX_FILES = 5
LINK_HOURS = 12
INLINE_TYPES = {"image/png", "image/jpeg", "image/webp", "image/gif", "application/pdf"}   # safe to show in a browser


def root() -> Path:
    path = Path(settings.uploads_dir)
    path.mkdir(parents=True, exist_ok=True)
    return path


def _ext(filename: str) -> str:
    suffix = Path(filename).suffix.lower()
    return suffix if re.fullmatch(r"\.[a-z0-9]{1,10}", suffix) else ""


def _clean_name(filename: str) -> str:
    name = Path(filename.replace("\\", "/")).name.replace("\x00", "").strip()
    return name[-200:] or "fayl"


async def save(user_id: int, upload: UploadFile) -> SubmissionFile:
    limit = settings.max_upload_mb * 1024 * 1024
    name = _clean_name(upload.filename or "")
    stored = f"{uuid.uuid4().hex}{_ext(name)}"
    path = root() / stored
    size = 0
    try:
        with path.open("wb") as out:
            while chunk := await upload.read(1024 * 1024):
                size += len(chunk)
                if size > limit:
                    raise HTTPException(413, f"Fayl {settings.max_upload_mb} MB dan katta bo'lmasin")
                out.write(chunk)
        if size == 0:
            raise HTTPException(422, "Fayl bo'sh")
    except BaseException:
        path.unlink(missing_ok=True)
        raise
    mime = mimetypes.guess_type(name)[0] or "application/octet-stream"
    return SubmissionFile(user_id=user_id, filename=name, stored_name=stored, mime=mime, size=size)


def file_out(f: SubmissionFile) -> dict:
    token = create_token(str(f.id), "file", LINK_HOURS)
    return {"id": f.id, "name": f.filename, "size": f.size, "mime": f.mime,
            "url": f"/api/v1/files/{f.id}/{quote(f.filename)}?t={token}"}


# --------------------------------------------------------------------------- reading files for the AI

AI_INLINE = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp",
             ".heic": "image/heic", ".pdf": "application/pdf"}
TEXT_EXTS = {".txt", ".csv", ".tsv", ".md", ".py", ".r", ".sql", ".do", ".json", ".js", ".ts", ".tsx", ".jsx",
             ".java", ".c", ".cpp", ".cs", ".go", ".php", ".html", ".css", ".xml", ".yaml", ".yml", ".log",
             ".tex", ".m", ".sas", ".dax", ".sps"}
TEXT_LIMIT = 40_000
INLINE_LIMIT = 14 * 1024 * 1024          # Gemini accepts about 20 MB inline per request


def _excel(path: Path) -> str:
    from openpyxl import load_workbook
    values = load_workbook(path, read_only=True, data_only=True)
    formulas = load_workbook(path, read_only=True, data_only=False)
    out = []
    try:
        for ws_v, ws_f in list(zip(values.worksheets, formulas.worksheets))[:8]:
            out.append(f"--- Varaq: {ws_v.title}")
            rows = zip(ws_v.iter_rows(max_row=200, max_col=30, values_only=True),
                       ws_f.iter_rows(max_row=200, max_col=30, values_only=True))
            for vals, forms in rows:
                cells = []
                for v, f in zip(vals, forms):
                    text = "" if v is None else str(v)
                    if isinstance(f, str) and f.startswith("="):
                        text = f"{text} [{f}]".strip()
                    cells.append(text)
                if any(cells):
                    out.append(" | ".join(cells).rstrip(" |"))
    finally:
        values.close()
        formulas.close()
    return "\n".join(out)


def _docx(path: Path) -> str:
    from docx import Document
    doc = Document(str(path))
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        lines += [" | ".join(c.text.strip() for c in row.cells) for row in table.rows]
    return "\n".join(lines)


def _pptx(path: Path) -> str:
    from pptx import Presentation
    lines = []
    for i, slide in enumerate(Presentation(str(path)).slides, 1):
        lines.append(f"--- Slayd {i}")
        lines += [s.text_frame.text for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
    return "\n".join(lines)


def _notebook(path: Path) -> str:
    nb = json.loads(path.read_text(encoding="utf-8", errors="replace"))
    lines = []
    for cell in nb.get("cells", []):
        lines.append(f"[{cell.get('cell_type')}]\n" + "".join(cell.get("source", [])))
        for o in cell.get("outputs", []):
            text = o.get("text") or o.get("data", {}).get("text/plain") or []
            if text:
                lines.append("[natija]\n" + "".join(text)[:2000])
    return "\n".join(lines)


def _as_text(path: Path, ext: str) -> str | None:
    if ext in TEXT_EXTS:
        return path.read_text(encoding="utf-8", errors="replace")
    reader = {".xlsx": _excel, ".xlsm": _excel, ".docx": _docx, ".pptx": _pptx, ".ipynb": _notebook}.get(ext)
    return reader(path) if reader else None


def ai_parts(attachments: list[tuple[str, str, int]]) -> list[dict]:
    """Gemini request parts for (original name, stored name, size) files. Blocking: run in a thread."""
    parts: list[dict] = []
    inline_total = 0
    for name, stored, size in attachments:
        path = root() / stored
        ext = _ext(name)
        if ext in AI_INLINE and inline_total + size <= INLINE_LIMIT:
            inline_total += size
            parts += [{"text": f"Fayl: {name}"},
                      {"inline_data": {"mime_type": AI_INLINE[ext], "data": base64.b64encode(path.read_bytes()).decode()}}]
            continue
        try:
            text = _as_text(path, ext)
        except Exception:                     # a broken or password-protected file: say so, do not fail
            text = None
        if text is None:
            parts.append({"text": f"Fayl: {name}\n(Bu fayl turini o'qib bo'lmadi, faqat nomi ma'lum.)"})
        else:
            cut = " (qisqartirildi)" if len(text) > TEXT_LIMIT else ""
            parts.append({"text": f"Fayl: {name}{cut}\n```\n{text[:TEXT_LIMIT]}\n```"})
    return parts
